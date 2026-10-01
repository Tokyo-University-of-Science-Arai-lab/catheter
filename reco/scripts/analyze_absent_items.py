#!/usr/bin/env python3
"""「リストにあるが画像に無い」「画像にあるがリストに無い」品目の、認識への影響を集計する。

背景(2026-09-26): 一対一対応(ハンガリー法)は「マスタの全品目が画像に映っている」前提では
成立しないという指摘があり、合計スコアS(5キー合計)によるしきい値足切りが使えるかを調べる。
reco/0911 は33品目が全部映っているため、以下の2つの状況を作って認識を回した。

  状況1「リストにあるが画像に無い」: 画像の左右10%を切り取った reco/0911_crop10 を、33品目のマスタで認識
                                        (make_cropped_dataset.py で作成)
  状況2「画像にあるがリストに無い」  : 元の reco/0911 を、10品目を削除した master_catheter_0911_2.json で認識
                                        (加えて、映っていないOPTIMA 2品目もマスタに足してある)

正解(どの箱がどの品目か)は、元の0911を9/11に認識し目視確認でTとなった行の「選択マスク」から取る
(--ref-result)。誤選択(F)の行は箱の位置が分からないので「不明」として集計から外す。

使い方(リポジトリルートから):
    python3 reco/scripts/analyze_absent_items.py \\
        --ref-result reco/0911/reco_result_20260911_172744 \\
        --crop-result reco/0911_crop10/reco_result_<日時> \\
        --unlisted-result reco/0911/reco_result_<日時> \\
        --out-dir <出力先>

出力(--out-dir): 状況1_行ごとの結果.csv / 状況2_行ごとの結果.csv / 集計.md
(足切りは実際には行わない。認識はフラグOFFで回した結果から、Sが基準未満なら「該当なし」になる
 ものとして、後から数え直している。MULTIKEY_REJECT_LOW_SCORE=1で回した場合と同じ判定になる。)
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import statistics
from pathlib import Path

import numpy as np
import openpyxl

THRESHOLDS = (100, 150, 200, 250, 300)
IOU_OK = 0.5            # 選択マスクと正解マスクのIoUがこれ以上なら「その箱を選んだ」
VIS_IN, VIS_OUT = 0.9, 0.1   # 切り取り範囲に入っている面積割合: >=0.9 映っている / <=0.1 映っていない / 間は見切れ


def load_run(result_dir: Path) -> list[dict]:
    """reco_result_*/work/N-*/ から、1ジョブ1行の辞書を作る(リトライ用の_attemptは除く)。"""
    rows = []
    for wd in sorted(glob.glob(str(result_dir / "work" / "*")), key=lambda p: int(os.path.basename(p).split("-")[0])):
        base = os.path.basename(wd)
        if "_attempt" in base:
            continue
        dbg_path = Path(wd) / "multikey_match_debug.json"
        if not dbg_path.exists():
            continue
        d = json.loads(dbg_path.read_text())
        S_by_mask = {p["mask"]: sum(p["by_key"].values()) for p in d["per_mask"]}
        sel = d["selected_mask"]
        bk = next((p["by_key"] for p in d["per_mask"] if p["mask"] == sel), None)
        rows.append(dict(n=int(base.split("-")[0]), shot=base.split("-")[1], query=d["query"], work=Path(wd),
                         sel=sel, S=S_by_mask.get(sel), S_by_mask=S_by_mask, by_key=bk, debug=d))
    return rows


def load_mask(work: Path, mask_name: str) -> np.ndarray:
    k = int(mask_name.split("_")[1])
    return np.load(work / "sam3_service_masks.npz")["masks"][k - 1]


def load_reference(ref_dir: Path) -> dict[tuple[str, str], np.ndarray]:
    """9/11の認識で目視確認Tだった行 -> {(画像番号, 品目): 正解マスク(元画像の座標)}"""
    xlsx = glob.glob(str(ref_dir / "catheter_width_report_*.xlsx"))[0]
    ws = openpyxl.load_workbook(xlsx, data_only=True)[[n for n in openpyxl.load_workbook(xlsx).sheetnames if n != "集計"][0]]
    hdr = [c.value for c in ws[1]]
    tf = {r[hdr.index("認識した順番")]: r[hdr.index("目視確認(T/F)")] for r in ws.iter_rows(min_row=2, values_only=True)}
    gt = {}
    for r in load_run(ref_dir):
        if tf.get(r["n"]) == "T":
            gt[(r["shot"], r["query"])] = load_mask(r["work"], r["sel"])
    return gt


def iou(a: np.ndarray, b: np.ndarray) -> float:
    u = np.logical_or(a, b).sum()
    return float(np.logical_and(a, b).sum() / u) if u else 0.0


def landing(pred: np.ndarray, shot: str, gt: dict, x0: int = 0, x1: int | None = None) -> tuple[str | None, float]:
    """選択マスクが、正解マスクのうちどの品目の箱と一番重なるか -> (品目, IoU)"""
    best, best_iou = None, 0.0
    for (s, item), m in gt.items():
        if s != shot:
            continue
        v = iou(pred, m[:, x0:x1])
        if v > best_iou:
            best, best_iou = item, v
    return (best if best_iou >= IOU_OK else None), best_iou


def stats(vals: list[float]) -> str:
    if not vals:
        return "-"
    v = sorted(vals)
    return f"n={len(v)} 最小{v[0]:.0f} / 25%点{v[len(v)//4]:.0f} / 中央値{statistics.median(v):.0f} / 最大{v[-1]:.0f}"


def sweep_table(groups: dict[str, list[float]], header: str) -> list[str]:
    """groups: 名前 -> Sのリスト。しきい値ごとに「S<しきい値=該当なし」になる件数を数える。"""
    lines = [f"| {header} | 件数 | " + " | ".join(f"S<{t}" for t in THRESHOLDS) + " |",
             "|---|---|" + "---|" * len(THRESHOLDS)]
    for name, vals in groups.items():
        cells = " | ".join(f"{sum(v < t for v in vals)}" for t in THRESHOLDS)
        lines.append(f"| {name} | {len(vals)} | {cells} |")
    return lines


# 「該当なし」にする条件の候補。by_key(選択マスクの5キーのスコア) -> True なら該当なし。
# 2026-09-26: 合計スコアSは、単位が読めない数値へ80点を与える等で、SPEC_1/SPEC_2/dateが
# 映っていない品目にも高得点を出す(状況1で、映っていない品目のspec_1は18件中13件が70点以上)。
# 識別力があるのは ref と display_name。この2つだけで判定する案を、Sと並べて比較する。
RULES = {
    "S(5キー合計)<200": lambda k: sum(k.values()) < 200,
    "ref+display_name<100": lambda k: k["ref"] + k["display_name"] < 100,
    "max(ref,display_name)<60": lambda k: max(k["ref"], k["display_name"]) < 60,
    "max(ref,display_name)<70": lambda k: max(k["ref"], k["display_name"]) < 70,
}


def rule_table(groups: dict[str, list[dict]]) -> list[str]:
    """groups: 区分名 -> 選択マスクのby_keyのリスト。各条件で「該当なし」になる件数を数える。"""
    lines = ["| 該当なしにする条件 | " + " | ".join(f"{n}({len(v)}件)" for n, v in groups.items()) + " |",
             "|---|" + "---|" * len(groups)]
    for name, f in RULES.items():
        lines.append(f"| {name} | " + " | ".join(f"{sum(map(f, v))}" for v in groups.values()) + " |")
    return lines


def write_csv(path: Path, rows: list[dict]) -> None:
    rows = [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def analyze_crop(gt, run, crop_info) -> tuple[list[dict], list[str]]:
    rows = []
    for r in run:
        cs = crop_info["shots"][r["shot"]]
        x0, x1 = cs["x0"], cs["x1"]
        g = gt.get((r["shot"], r["query"]))
        pred = load_mask(r["work"], r["sel"])
        if g is None:
            cls, vis, ok, lands = "不明(正解の位置が分からない)", None, None, None
        else:
            vis = float(g[:, x0:x1].sum() / g.sum())
            cls = "映っている" if vis >= VIS_IN else ("映っていない" if vis <= VIS_OUT else "見切れ")
            v = iou(pred, g[:, x0:x1])
            ok = v >= IOU_OK
            lands, _ = landing(pred, r["shot"], gt, x0, x1)
        rows.append(dict(画像=r["shot"], 品目=r["query"], 区分=cls, 映っている割合=None if vis is None else round(vis, 2),
                         S=round(r["S"], 1), ref=r["by_key"]["ref"], display_name=r["by_key"]["display_name"], date=r["by_key"]["date"],
                         spec_1=r["by_key"]["spec_1"], spec_2=r["by_key"]["spec_2"],
                         選択マスク=r["sel"], 正解を選んだ=ok, 選択マスクが重なる品目=lands, 順番=r["n"], _bk=r["by_key"]))
    S = lambda cond: [x["S"] for x in rows if cond(x)]
    g_ok = S(lambda x: x["区分"] == "映っている" and x["正解を選んだ"])
    g_ng = S(lambda x: x["区分"] == "映っている" and x["正解を選んだ"] is False)
    g_out = S(lambda x: x["区分"] == "映っていない")
    g_cut = S(lambda x: x["区分"] == "見切れ")
    n_by = {c: sum(x["区分"] == c for x in rows) for c in ("映っている", "見切れ", "映っていない")}
    md = ["## 状況1: リストにあるが画像に無い品目(reco/0911_crop10、33品目のマスタ)", "",
          f"- 切り取り: 左右{crop_info['ratio']*100:.0f}%ずつ(x=0〜{crop_info['shots']['1']['x0']}と{crop_info['shots']['1']['x1']}〜{crop_info['shots']['1']['元の幅']}を除去)",
          f"- 区分の定義: 正解の箱が切り取り範囲に入っている面積割合が {VIS_IN:.0%}以上=映っている / {VIS_OUT:.0%}以下=映っていない / 間=見切れ",
          f"- 件数: 映っている {n_by['映っている']} / 見切れ {n_by['見切れ']} / **映っていない {n_by['映っていない']}** / 不明 {sum(x['区分'].startswith('不明') for x in rows)}"
          f" (計{len(rows)})", "",
          "### 映っている品目の正解率(参考: 元の0911では、Tの行だけを使うため常に100%)", "",
          f"- 正解を選んだ: {len(g_ok)}件 / 誤選択: {len(g_ng)}件 (正解率 {len(g_ok) / max(1, len(g_ok) + len(g_ng)):.1%})", "",
          "### 選択マスクの合計スコアS", "",
          f"- 映っている・正解: {stats(g_ok)}", f"- 映っている・誤選択: {stats(g_ng)}",
          f"- **映っていない(強制的に選ばれた箱)**: {stats(g_out)}", f"- 見切れ(参考): {stats(g_cut)}", "",
          "### しきい値ごとに「該当なし」になる件数(S<しきい値)", ""]
    md += sweep_table({"映っている・正解(なるべく落としたくない)": g_ok, "映っている・誤選択": g_ng,
                       "映っていない(該当なしにしたい)": g_out, "見切れ(参考)": g_cut}, "区分")
    md += ["", "### 該当なしにする条件の比較(Sの合計と、識別力のあるref/display_nameだけを使う案)", ""]
    md += rule_table({"映っている・正解": [x["_bk"] for x in rows if x["区分"] == "映っている" and x["正解を選んだ"]],
                      "映っている・誤選択": [x["_bk"] for x in rows if x["区分"] == "映っている" and x["正解を選んだ"] is False],
                      "映っていない": [x["_bk"] for x in rows if x["区分"] == "映っていない"]})
    landed = [x for x in rows if x["区分"] == "映っていない"]
    if landed:
        from collections import Counter
        c = Counter(("映っている品目の箱を奪った" if x["選択マスクが重なる品目"] else "どの正解箱とも重ならない") for x in landed)
        md += ["", "### 映っていない品目が、強制的に選んだ箱", "",
               f"- {dict(c)} (正解箱が分かっている品目の箱に重なれば、その品目から箱を奪っている)"]
    return rows, md


def analyze_unlisted(gt, run, master_items, all_master_items) -> tuple[list[dict], list[str]]:
    removed = set(all_master_items) - set(master_items)     # 画像にあるがリストに無い品目
    rows = []
    for r in run:
        g = gt.get((r["shot"], r["query"]))
        pred = load_mask(r["work"], r["sel"])
        lands, lv = landing(pred, r["shot"], gt)
        if r["query"] not in all_master_items:
            cls = "リストにあるが画像に無い(追加した品目)"
        elif g is None:
            cls = "不明(正解の位置が分からない)"
        else:
            cls = "リストにも画像にもある"
        ok = None if g is None or cls.startswith("リストにあるが") else iou(pred, g) >= IOU_OK
        if lands is None:
            where = "どの正解箱とも重ならない"
        elif lands == r["query"]:
            where = "自分の箱(正解)"
        elif lands in removed:
            where = "リストに無い品目の箱を奪った"
        else:
            where = "別のリスト品目の箱を奪った"
        rows.append(dict(画像=r["shot"], 品目=r["query"], 区分=cls, S=round(r["S"], 1), ref=r["by_key"]["ref"],
                         display_name=r["by_key"]["display_name"], date=r["by_key"]["date"], spec_1=r["by_key"]["spec_1"],
                         spec_2=r["by_key"]["spec_2"], 選択マスク=r["sel"], 正解を選んだ=ok, 選択マスクの行き先=where,
                         順番=r["n"], _bk=r["by_key"]))
    S = lambda cond: [x["S"] for x in rows if cond(x)]
    both = [x for x in rows if x["区分"] == "リストにも画像にもある"]
    g_ok = S(lambda x: x["区分"] == "リストにも画像にもある" and x["正解を選んだ"])
    g_ng = S(lambda x: x["区分"] == "リストにも画像にもある" and x["正解を選んだ"] is False)
    g_add = S(lambda x: x["区分"].startswith("リストにあるが"))
    from collections import Counter
    wrong_where = Counter(x["選択マスクの行き先"] for x in both if x["正解を選んだ"] is False)
    add_where = Counter(x["選択マスクの行き先"] for x in rows if x["区分"].startswith("リストにあるが"))
    md = ["## 状況2: 画像にあるがリストに無い品目(reco/0911 の画像 + 25品目のマスタ)", "",
          f"- リストから外した品目({len(removed)}件): {', '.join(sorted(removed))}",
          f"- マスタに足した「映っていない」品目: {', '.join(sorted(set(master_items) - set(all_master_items)))}(0911の写真には映っていないと判断)",
          f"- 件数: リストにも画像にもある(正解の位置が分かる) {len(both)} / 追加した品目 {len(g_add)}"
          f" / 不明 {sum(x['区分'].startswith('不明') for x in rows)} (計{len(rows)})", "",
          "### リストにも画像にもある品目の正解率(元の0911ではTの行だけなので、常に100%が基準)", "",
          f"- 正解を選んだ: {len(g_ok)}件 / 誤選択: {len(g_ng)}件 (正解率 {len(g_ok) / max(1, len(both)):.1%})",
          f"- 誤選択の行き先: {dict(wrong_where)}", "",
          "### 選択マスクの合計スコアS", "",
          f"- リストにも画像にもある・正解: {stats(g_ok)}", f"- リストにも画像にもある・誤選択: {stats(g_ng)}",
          f"- **追加した品目(映っていない)**: {stats(g_add)}",
          f"- 追加した品目が強制的に選んだ箱: {dict(add_where)}", "",
          "### しきい値ごとに「該当なし」になる件数(S<しきい値)", ""]
    md += sweep_table({"リストにも画像にもある・正解(なるべく落としたくない)": g_ok, "リストにも画像にもある・誤選択": g_ng,
                       "追加した品目(該当なしにしたい)": g_add}, "区分")
    md += ["", "### 該当なしにする条件の比較(Sの合計と、識別力のあるref/display_nameだけを使う案)", ""]
    md += rule_table({"リストにも画像にもある・正解": [x["_bk"] for x in rows if x["区分"] == "リストにも画像にもある" and x["正解を選んだ"]],
                      "リストにも画像にもある・誤選択": [x["_bk"] for x in rows if x["区分"] == "リストにも画像にもある" and x["正解を選んだ"] is False],
                      "追加した品目(映っていない)": [x["_bk"] for x in rows if x["区分"].startswith("リストにあるが")]})
    return rows, md


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ref-result", required=True, help="正解の位置を取る、元0911の認識結果(目視確認済み)")
    ap.add_argument("--crop-result", help="状況1: 切り取り版の認識結果 reco_result_*")
    ap.add_argument("--unlisted-result", help="状況2: 元画像+削除したマスタでの認識結果 reco_result_*")
    ap.add_argument("--full-master", default="Master_JSON/master_catheter_0911.json", help="元の33品目のマスタ")
    ap.add_argument("--out-dir", required=True)
    args = ap.parse_args()

    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    gt = load_reference(Path(args.ref_result))
    print(f"正解の位置が分かる(画像, 品目): {len(gt)}件")
    md = ["# 「該当なし」足切りの検証: 状況1と状況2", "",
          f"正解の位置の出典: {args.ref_result}(目視確認T、{len(gt)}件)", ""]

    if args.crop_result:
        crop_dir = Path(args.crop_result)
        crop_info = json.loads((crop_dir.parent / "crop_info.json").read_text())
        rows, m = analyze_crop(gt, load_run(crop_dir), crop_info)
        write_csv(out / "状況1_行ごとの結果.csv", rows)
        md += m + [""]
    if args.unlisted_result:
        run = load_run(Path(args.unlisted_result))
        master_items = [x["book_name"] for x in json.loads(Path(run[0]["debug"]["master_json"]).read_text())]
        all_items = {x["book_name"] for x in json.loads(Path(args.full_master).read_text())}
        rows, m = analyze_unlisted(gt, run, set(master_items), all_items)
        write_csv(out / "状況2_行ごとの結果.csv", rows)
        md += m + [""]
    (out / "集計.md").write_text("\n".join(md), encoding="utf-8")
    print("\n".join(md))


if __name__ == "__main__":
    main()
